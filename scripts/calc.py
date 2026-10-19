"""Расчёт оклада и допов по КПП и журналу фактических дней."""
import glob, json, re, sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
OFF = ("ВЫХОДНОЙ", "ОТСЫПНОЙ", "ПОДГОТОВ")


def parse_kpp(path):
    """Возвращает {ДД.ММ: статус} по строкам-заголовкам дней КПП."""
    ws = openpyxl.load_workbook(path).worksheets[0]
    days = {}
    for row in ws.iter_rows(values_only=True):
        m = re.match(r"(\d\d\.\d\d)\.\d{4}$", str(row[0] or ""))
        if not m:
            continue
        text = " ".join(str(c) for c in row[1:6] if c is not None)
        sm = re.search(r"СМЕНА\s*№\s*([\d-]+)", text)
        if sm:
            days[m[1]] = ("shift", f"СМЕНА №{sm[1]} ✅")
        else:
            off = next((o for o in OFF if o in text.upper()), None)
            days[m[1]] = ("off", off + ("КА" if off == "ПОДГОТОВ" else "")) if off else ("none", "режим не указан")
    return days


def key(d):
    return (d[3:], d[:2])


def main():
    kpp_file = sorted(glob.glob(str(ROOT / "kpp" / "*.xlsx")))[-1]
    kpp = parse_kpp(kpp_file)
    j = json.loads((ROOT / "data" / "journal.json").read_text(encoding="utf-8"))
    actual = {d["date"]: d["action"] for d in j["days"]}

    rows, warn = [], []
    for d in sorted(set(kpp) | set(actual), key=key):
        kind, status = kpp.get(d, ("outside", "Вне КПП (тех. день)"))
        act = actual.get(d)
        if kind == "shift":
            extra, action = 0, act or "Смена по КПП (в окладе)"
            if act:
                warn.append(f"{d}: в журнале работа, но по КПП это съёмочная смена — допа нет")
        elif kind == "outside":
            extra, action = j["extra_rate"], act
        elif act:
            extra, action = j["extra_rate"], act
        else:
            extra, action = 0, "Нет работы"
        rows.append((d, status, action, extra))
    shifts = sum(1 for v in kpp.values() if v[0] == "shift")
    extras = sum(r[3] for r in rows)
    total = j["base_salary"] + extras

    print(f"КПП: {Path(kpp_file).name}")
    print(f"{'Дата':6} | {'Статус по КПП':34} | {'Фактическое действие':50} | {'Допл.':>6} | Итого")
    run = j["base_salary"]
    out = []
    for d, s, a, x in rows:
        run += x
        out.append((d, s, a, x, run))
        print(f"{d:6} | {s:34} | {a[:50]:50} | {x:>6} | {run}")
    print(f"\nСмен в окладе по КПП: {shifts}")
    print(f"Оклад {j['base_salary']} + допы {extras} = {total}")
    for w in warn:
        print("ВНИМАНИЕ:", w)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчет по сменам"
    ws.append(["Отчет по сменам — СУПЕР-ВОВИК, " + Path(kpp_file).name])
    ws.append([])
    ws.append(["Дата", "Плановый статус по КПП", "Фактическое действие", "Доплата, ₽", "Накопительно, ₽"])
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for r in out:
        ws.append(list(r))
        if r[3]:
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor="FFF2CC")
    ws.append([])
    ws.append(["Оклад", None, None, None, j["base_salary"]])
    ws.append([f"Допы", None, None, None, extras])
    ws.append(["К выплате", None, None, None, total])
    ws.append([f"Смен в окладе по КПП: {shifts}"])
    for col, w in zip("ABCDE", (10, 36, 52, 14, 18)):
        ws.column_dimensions[col].width = w
    ws["A1"].font = Font(bold=True, size=13)
    (ROOT / "reports").mkdir(exist_ok=True)
    wb.save(ROOT / "reports" / "otchet_po_smenam.xlsx")


if __name__ == "__main__":
    sys.exit(main())
